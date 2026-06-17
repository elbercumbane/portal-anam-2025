#!/usr/bin/env python3
"""
Importa membros a partir das fichas Word (ANAM_Ficha_Inscricao_*.docx)
e gera script SQL de backup.
"""

from __future__ import annotations

import argparse
import glob
import os
import re
import sys
import unicodedata
import uuid
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("Instale python-docx: pip install python-docx")
    sys.exit(1)

try:
    import psycopg2
    from psycopg2.extras import execute_values
except ImportError:
    psycopg2 = None

BASE_DIR = Path(__file__).resolve().parents[2]
DEFAULT_DOCX_DIR = BASE_DIR / "dados dos membros"
DEFAULT_SQL_OUT = BASE_DIR / "backend" / "sql" / "import-membros-anam-backup.sql"

DB_CONFIG = {
    "host": os.getenv("DB_HOST", "localhost"),
    "port": int(os.getenv("DB_PORT", "5432")),
    "dbname": os.getenv("DB_NAME", "anam"),
    "user": os.getenv("DB_USER", "postgres"),
    "password": os.getenv("DB_PASSWORD", "admin123"),
}

DOTS_RE = re.compile(r"\.{3,}")
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:\+?258)?\s*8[2-7]\d{7}(?:\s*/\s*(?:\+?258)?\s*8[2-7]\d{7})?")
DATE_RE = re.compile(r"(\d{1,2})\s*/\s*(\d{1,2})\s*/\s*(\d{2,4})")


def clean_text(value: str | None) -> str:
    if not value:
        return ""
    value = value.replace("\xa0", " ").replace("\r", "\n")
    value = DOTS_RE.sub("", value)
    value = re.sub(r"\s+", " ", value).strip()
    return value


def slugify(text: str) -> str:
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = re.sub(r"[^a-zA-Z0-9]+", ".", text.lower()).strip(".")
    return text or "membro"


def parse_date(raw: str | None) -> str | None:
    if not raw:
        return None
    match = DATE_RE.search(raw)
    if not match:
        return None
    day, month, year = match.groups()
    if not day.strip() or not month.strip() or not year.strip():
        return None
    day_i = int(day)
    month_i = int(month)
    year_i = int(year)
    if year_i < 100:
        year_i += 1900 if year_i >= 30 else 2000
    try:
        return f"{year_i:04d}-{month_i:02d}-{day_i:02d}"
    except ValueError:
        return None


def extract_after_label(text: str, labels: list[str]) -> str:
    for label in labels:
        pattern = re.compile(rf"{re.escape(label)}\s*:?\s*(.*)$", re.IGNORECASE | re.DOTALL)
        match = pattern.search(text)
        if match:
            value = clean_text(match.group(1))
            if value and not value.startswith("_"):
                return value.split("\n")[0].strip()
    return ""


def parse_estado_civil(line: str) -> str:
    line_lower = line.lower()
    options = [
        ("solteiro", "Solteiro(a)"),
        ("casado", "Casado(a)"),
        ("união marital", "União Marital"),
        ("uniao marital", "União Marital"),
        ("união de facto", "União de Facto"),
        ("viúvo", "Viúvo(a)"),
        ("viuvo", "Viúvo(a)"),
        ("divorciado", "Divorciado(a)"),
    ]
    for key, label in options:
        pattern = re.compile(rf"{re.escape(key)}[^a-zA-Z]*x", re.IGNORECASE)
        if pattern.search(line_lower) or re.search(rf"x\s*{re.escape(key)}", line_lower):
            return label
    for key, label in options:
        if key in line_lower:
            return label
    return "Não informado"


def parse_address(line: str) -> tuple[str, str, str]:
    rua = bairro = cidade = ""
    if not line:
        return "Não informado", "Não informado", "Maputo"

    rua_match = re.search(r"Rua:\s*(.*?)(?:;|$)", line, re.IGNORECASE)
    bairro_match = re.search(r"Bairro\s+([^;]+)", line, re.IGNORECASE)
    if rua_match:
        rua = clean_text(rua_match.group(1)) or "Não informado"
    if bairro_match:
        bairro = clean_text(bairro_match.group(1)) or "Não informado"

    city_match = re.search(r"(Maputo|Massinga|Beira|Nampula|Quelimane|Inhambane|Xai-Xai|Pemba)", line, re.IGNORECASE)
    cidade = city_match.group(1) if city_match else "Maputo"
    return rua or "Não informado", bairro or "Não informado", cidade


def extract_phones(text: str) -> tuple[str, str | None]:
    phones = []
    for match in PHONE_RE.finditer(text):
        phones.extend([p.strip() for p in match.group(0).split("/") if p.strip()])
    phones = [re.sub(r"\s+", "", p) for p in phones if p]
    if not phones:
        return "000000000", None
    celular = phones[0].replace("+258", "").replace(" ", "")
    telefone = phones[1].replace("+258", "").replace(" ", "") if len(phones) > 1 else None
    return celular[:20], (telefone[:20] if telefone else None)


def normalize_email(email: str) -> str:
    email = email.lower().strip()
    email = re.sub(r"\.+", ".", email)
    email = email.replace(".@", "@")
    email = email.strip(".")
    return email[:255]


def extract_email(text: str, fallback_name: str, used_emails: set[str]) -> str:
    emails = [normalize_email(e) for e in EMAIL_RE.findall(text)]
    emails = [e for e in emails if e and "@" in e and not e.endswith(".")]
    for email in emails:
        if email not in used_emails:
            used_emails.add(email)
            return email
    base = slugify(fallback_name)
    candidate = f"{base}@anam-import.mz"
    index = 1
    while candidate in used_emails:
        candidate = f"{base}.{index}@anam-import.mz"
        index += 1
    used_emails.add(candidate)
    return candidate


def parse_docx(path: Path, used_emails: set[str]) -> dict:
    doc = Document(str(path))
    paras = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    full_text = "\n".join(paras)

    nome = ""
    nome_pai = ""
    nome_mae = ""
    estado_civil = "Não informado"
    nome_conjuge = None
    data_nasc = None
    data_nasc_conjuge = None
    filiacao_pai_conjuge = None
    filiacao_mae_conjuge = None
    profissao = "Não informado"
    local_trabalho = None
    endereco_trabalho = None
    telefone_trabalho = None
    outras = None

    for i, para in enumerate(paras):
        if para.startswith("Nome") and "\n" in para and not nome:
            nome = clean_text(para.split("\n", 1)[1])
        elif para.startswith("Filho de") and "\n" in para and not nome_pai:
            nome_pai = clean_text(para.split("\n", 1)[1])
        elif para.startswith("E de") and "\n" in para and not nome_mae:
            nome_mae = clean_text(para.split("\n", 1)[1])
        elif para.lower().startswith("estado civil"):
            estado_civil = parse_estado_civil(para)
        elif "conjug" in para.lower() and "\n" in para:
            nome_conjuge = clean_text(para.split("\n", 1)[1]) or None
        elif para.lower().startswith("filho/a de") and "\n" in para:
            filiacao_pai_conjuge = clean_text(para.split("\n", 1)[1]) or None
        elif para.lower().startswith("e de") and "\n" in para and filiacao_pai_conjuge and not filiacao_mae_conjuge:
            filiacao_mae_conjuge = clean_text(para.split("\n", 1)[1]) or None
        elif para.lower().startswith("profiss"):
            if "\n" in para:
                profissao = clean_text(para.split("\n", 1)[1]) or profissao
            elif i + 1 < len(paras):
                profissao = clean_text(paras[i + 1]) or profissao
        elif para.lower().startswith("local de trabalho") and i + 1 < len(paras):
            nxt = clean_text(paras[i + 1])
            if nxt and not nxt.lower().startswith("profiss") and not set(nxt) <= {"_"}:
                local_trabalho = nxt
        elif para.lower().startswith("endereço:") or para.lower().startswith("endereco:"):
            endereco_trabalho = clean_text(para.split(":", 1)[-1]) or None

    birth_dates = [parse_date(p) for p in paras if "data de nascimento" in p.lower()]
    birth_dates = [d for d in birth_dates if d]
    if birth_dates:
        data_nasc = birth_dates[0]
    if len(birth_dates) > 1:
        data_nasc_conjuge = birth_dates[1]

    if not nome:
        stem = path.stem.replace("ANAM_Ficha_Inscricao_", "").replace("_", " ")
        nome = clean_text(stem)

    address_line = next((p for p in paras if p.lower().startswith("rua:")), "")
    rua, bairro, cidade = parse_address(address_line)
    celular, telefone = extract_phones(full_text)
    email = extract_email(full_text, nome, used_emails)

    if not nome_pai:
        nome_pai = "Não informado"
    if not nome_mae:
        nome_mae = "Não informado"
    if not data_nasc:
        data_nasc = "1970-01-01"

    work_phone_match = re.search(r"Tel:\s*([^;]+).*?Cel:\s*([^;]+)", full_text, re.IGNORECASE)
    if work_phone_match:
        tel = clean_text(work_phone_match.group(1))
        cel = clean_text(work_phone_match.group(2))
        parts = [p for p in [tel, cel] if p and "." not in p]
        if parts:
            telefone_trabalho = " / ".join(parts)[:20]

    commitment = next((p for p in paras if "compromisso de honra" in p.lower()), "")
    if commitment:
        outras = clean_text(commitment)[:2000] or None

    return {
        "id": str(uuid.uuid4()),
        "foto": None,
        "nome_completo": nome[:255],
        "nome_pai": nome_pai[:255],
        "nome_mae": nome_mae[:255],
        "data_nascimento": data_nasc,
        "estado_civil": estado_civil[:50],
        "profissao": (profissao or "Não informado")[:255],
        "celular": celular[:20],
        "email": email[:255],
        "telefone": telefone,
        "rua": rua[:255],
        "bairro": bairro[:255],
        "cidade": cidade[:255],
        "nome_conjuge": (nome_conjuge[:255] if nome_conjuge else None),
        "data_nascimento_conjuge": data_nasc_conjuge,
        "filiacao_pai_conjuge": (filiacao_pai_conjuge[:255] if filiacao_pai_conjuge else None),
        "filiacao_mae_conjuge": (filiacao_mae_conjuge[:255] if filiacao_mae_conjuge else None),
        "local_trabalho": (local_trabalho[:255] if local_trabalho else None),
        "cargo": None,
        "endereco_trabalho": (endereco_trabalho[:255] if endereco_trabalho else None),
        "telefone_trabalho": telefone_trabalho,
        "outras_informacoes": outras,
        "aceita_termos": True,
        "status": "Aprovado",
        "data_cadastro": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "data_aprovacao": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "observacoes": f"Importado automaticamente de {path.name}",
        "source_file": path.name,
    }


def sql_literal(value) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    escaped = str(value).replace("'", "''")
    return f"'{escaped}'"


def generate_sql(members: list[dict], output_path: Path) -> None:
    lines = [
        "-- =====================================================",
        "-- BACKUP / IMPORTAÇÃO DE MEMBROS ANAM",
        f"-- Gerado em: {datetime.now().isoformat(timespec='seconds')}",
        f"-- Total de membros: {len(members)}",
        "-- Execute na base de dados 'anam'",
        "-- =====================================================",
        "",
        "CREATE EXTENSION IF NOT EXISTS \"uuid-ossp\";",
        "",
        "-- Remove membros importados anteriormente (opcional)",
        "-- DELETE FROM membros WHERE observacoes LIKE 'Importado automaticamente de ANAM_Ficha_Inscricao_%';",
        "",
    ]

    for member in members:
        cols = [
            "id", "foto", "nome_completo", "nome_pai", "nome_mae", "data_nascimento",
            "estado_civil", "profissao", "celular", "email", "telefone", "rua", "bairro", "cidade",
            "nome_conjuge", "data_nascimento_conjuge", "filiacao_pai_conjuge", "filiacao_mae_conjuge",
            "local_trabalho", "cargo", "endereco_trabalho", "telefone_trabalho", "outras_informacoes",
            "aceita_termos", "status", "data_cadastro", "data_aprovacao", "observacoes",
        ]
        values = ", ".join(sql_literal(member.get(col)) for col in cols)
        lines.append(
            "INSERT INTO membros (" + ", ".join(cols) + ") VALUES (" + values + ")\n"
            "ON CONFLICT (email) DO UPDATE SET\n"
            "  nome_completo = EXCLUDED.nome_completo,\n"
            "  nome_pai = EXCLUDED.nome_pai,\n"
            "  nome_mae = EXCLUDED.nome_mae,\n"
            "  data_nascimento = EXCLUDED.data_nascimento,\n"
            "  estado_civil = EXCLUDED.estado_civil,\n"
            "  profissao = EXCLUDED.profissao,\n"
            "  celular = EXCLUDED.celular,\n"
            "  telefone = EXCLUDED.telefone,\n"
            "  rua = EXCLUDED.rua,\n"
            "  bairro = EXCLUDED.bairro,\n"
            "  cidade = EXCLUDED.cidade,\n"
            "  nome_conjuge = EXCLUDED.nome_conjuge,\n"
            "  data_nascimento_conjuge = EXCLUDED.data_nascimento_conjuge,\n"
            "  filiacao_pai_conjuge = EXCLUDED.filiacao_pai_conjuge,\n"
            "  filiacao_mae_conjuge = EXCLUDED.filiacao_mae_conjuge,\n"
            "  local_trabalho = EXCLUDED.local_trabalho,\n"
            "  cargo = EXCLUDED.cargo,\n"
            "  endereco_trabalho = EXCLUDED.endereco_trabalho,\n"
            "  telefone_trabalho = EXCLUDED.telefone_trabalho,\n"
            "  outras_informacoes = EXCLUDED.outras_informacoes,\n"
            "  aceita_termos = EXCLUDED.aceita_termos,\n"
            "  status = EXCLUDED.status,\n"
            "  data_aprovacao = EXCLUDED.data_aprovacao,\n"
            "  observacoes = EXCLUDED.observacoes;"
        )
        lines.append("")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(lines), encoding="utf-8")


def insert_members_db(members: list[dict]) -> tuple[int, int]:
    if psycopg2 is None:
        raise RuntimeError("psycopg2 não instalado")

    inserted = 0
    updated = 0
    conn = psycopg2.connect(**DB_CONFIG)
    try:
        with conn.cursor() as cur:
            for member in members:
                cur.execute("SELECT id FROM membros WHERE email = %s", (member["email"],))
                existing = cur.fetchone()
                cols = [
                    "id", "foto", "nome_completo", "nome_pai", "nome_mae", "data_nascimento",
                    "estado_civil", "profissao", "celular", "email", "telefone", "rua", "bairro", "cidade",
                    "nome_conjuge", "data_nascimento_conjuge", "filiacao_pai_conjuge", "filiacao_mae_conjuge",
                    "local_trabalho", "cargo", "endereco_trabalho", "telefone_trabalho", "outras_informacoes",
                    "aceita_termos", "status", "data_cadastro", "data_aprovacao", "observacoes",
                ]
                values = [member.get(c) for c in cols]
                if existing:
                    set_clause = ", ".join(f"{c} = %s" for c in cols if c not in ("id", "email", "data_cadastro"))
                    update_vals = [member.get(c) for c in cols if c not in ("id", "email", "data_cadastro")]
                    update_vals.append(member["email"])
                    cur.execute(
                        f"UPDATE membros SET {set_clause} WHERE email = %s",
                        update_vals,
                    )
                    updated += 1
                else:
                    placeholders = ", ".join(["%s"] * len(cols))
                    cur.execute(
                        f"INSERT INTO membros ({', '.join(cols)}) VALUES ({placeholders})",
                        values,
                    )
                    inserted += 1
        conn.commit()
    finally:
        conn.close()
    return inserted, updated


def main() -> int:
    parser = argparse.ArgumentParser(description="Importar membros ANAM a partir de ficheiros Word")
    parser.add_argument("--docx-dir", default=str(DEFAULT_DOCX_DIR))
    parser.add_argument("--sql-out", default=str(DEFAULT_SQL_OUT))
    parser.add_argument("--dry-run", action="store_true", help="Apenas gera SQL, não insere na BD")
    parser.add_argument("--no-db", action="store_true", help="Não inserir na base de dados")
    args = parser.parse_args()

    docx_dir = Path(args.docx_dir)
    files = sorted(docx_dir.glob("ANAM_Ficha_Inscricao_*.docx"))
    if not files:
        print(f"Nenhum ficheiro encontrado em: {docx_dir}")
        return 1

    used_emails: set[str] = set()
    members: list[dict] = []
    errors: list[str] = []

    for file_path in files:
        try:
            member = parse_docx(file_path, used_emails)
            members.append(member)
            print(f"OK  {file_path.name} -> {member['nome_completo']} | {member['email']}")
        except Exception as exc:
            errors.append(f"{file_path.name}: {exc}")
            print(f"ERR {file_path.name}: {exc}")

    sql_out = Path(args.sql_out)
    generate_sql(members, sql_out)
    print(f"\nSQL gerado: {sql_out} ({len(members)} membros)")

    if not args.dry_run and not args.no_db:
        try:
            inserted, updated = insert_members_db(members)
            print(f"Base de dados: {inserted} inseridos, {updated} atualizados")
        except Exception as exc:
            print(f"Erro ao inserir na BD: {exc}")
            print("O ficheiro SQL foi gerado e pode ser executado manualmente.")
            return 2

    if errors:
        print(f"\n{len(errors)} ficheiro(s) com erro.")
        return 3

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
